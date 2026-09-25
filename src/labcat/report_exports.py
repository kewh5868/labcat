"""Download persisted reports without re-running research or trusting
markup.

ReportLab:
https://docs.reportlab.com/reportlab/userguide/ch6_paragraphs/
 python-docx:
https://python-docx.readthedocs.io/en/latest/user/quickstart.html
"""

import io
import json
import re
import secrets
import sqlite3
import textwrap
import threading
import time
from collections import OrderedDict
from copy import deepcopy
from dataclasses import asdict
from html import escape
from importlib.resources import files
from typing import Literal

from labcat.chemical_names import (
    ChemicalNameRequest,
    ChemicalNameResolver,
    offline_names,
)
from labcat.config import (
    ReportLayout,
    config_from_mapping,
    load_config,
    validate_layout,
)
from labcat.science.formula_display import display_formula, flat_formula_tokens
from labcat.science.ranking import score_color
from labcat.science.reporting import (
    PRESENTATION_VERSION,
    _citation_title,
    _preliminary_literature_lines,
    approved_citation_url,
    citation_references,
    render_reports,
    screening_priority_presentation,
)
from labcat.workspace import WorkspaceNotFound

FORMATS = {"text", "json", "pdf", "docx"}
VIEWS = {
    "pi": ("pi",),
    "audit": ("audit",),
    "both": ("pi", "audit"),
    "sources": ("sources",),
    "pi,sources": ("pi", "sources"),
    "audit,sources": ("audit", "sources"),
    "all": ("pi", "audit", "sources"),
}
VIEW_FIELDS = {"pi": "pi_summary", "audit": "technical_audit"}
VIEW_LABELS = {"pi": "Summary", "audit": "Technical View", "sources": "Sources"}
MAX_REPORT_CHARACTERS = 2_000_000
MAX_TABLE_ROWS = 1000
MAX_TABLE_COLUMNS = 103
MAX_TABLE_CELL_CHARACTERS = 4096
URL_PATTERN = re.compile(r"https://[^\s<>\"']+")
SOURCE_HOSTS = {
    "doi.org",
    "www.nature.com",
    "materialsproject.org",
    "api.materialsproject.org",
    "docs.materialsproject.org",
    "datadryad.org",
    "api.figshare.com",
    "ndownloader.figshare.com",
    "materials.hybrid3.duke.edu",
    "nomad-lab.eu",
    "europepmc.org",
    "arxiv.org",
}
_FONT_LOCK = threading.Lock()
ACCENTS = {
    "sage": ("#365B48", "#E2EEE7", "#F5F8F6"),
    "teal": ("#174B4B", "#DDEEEF", "#F3F8F8"),
    "slate": ("#35465B", "#E5EAF0", "#F6F7F9"),
}


class ExportError(ValueError):
    """A safe export validation error, without private paths or report
    contents."""


class ExportUnavailable(RuntimeError):
    """An optional document renderer is unavailable."""


def report_layout(report: dict, fallback: ReportLayout | None = None):
    """Saved report appearance wins; legacy reports explicitly use
    current settings."""
    result = report.get("result")
    execution = result.get("execution") if isinstance(result, dict) else None
    presentation = (
        execution.get("presentation") if isinstance(execution, dict) else None
    )
    if isinstance(presentation, dict) and "layout" in presentation:
        try:
            return validate_layout(presentation["layout"]), "saved-report"
        except ValueError as error:
            raise ExportError("The saved report appearance is invalid.") from error
    if fallback is not None:
        return validate_layout(asdict(fallback)), "current-settings"
    return ReportLayout(), "defaults"


def _parsed(text: str):
    """Keep JSON report strings as objects, avoiding escaped JSON inside
    JSON."""
    try:
        value = json.loads(text)
        return value if isinstance(value, (dict, list)) else text
    except (ValueError, TypeError, RecursionError):
        return text


def _content(report: dict, views: str) -> tuple[dict, dict]:
    if not isinstance(report, dict) or views not in VIEWS:
        raise ExportError("Choose at least one of Summary, Technical View, or Sources.")
    # Persisted reports carry exact workspace source memberships. Direct exports
    # obey the same boundary as the HTTP route, which already filters these IDs.
    if "source_ids" in report:
        source_ids, sources = report["source_ids"], report.get("sources", [])
        if (
            not isinstance(source_ids, list)
            or len(source_ids) > 1000
            or any(
                not isinstance(identity, str) or not 1 <= len(identity) <= 128
                for identity in source_ids
            )
            or not isinstance(sources, list)
            or len(sources) > 1000
            or any(
                not isinstance(source, dict) or not isinstance(source.get("id"), str)
                for source in sources
            )
            or set(source_ids) != {source.get("id") for source in sources}
            or any(
                (
                    "report_ids" in source
                    and (
                        not isinstance(source["report_ids"], list)
                        or report.get("id") not in source["report_ids"]
                    )
                )
                or (
                    "chat_ids" in source
                    and "chat_id" in report
                    and (
                        not isinstance(source["chat_ids"], list)
                        or report["chat_id"] not in source["chat_ids"]
                    )
                )
                for source in sources
            )
        ):
            raise ExportError(
                "The saved scientific report cannot be formatted safely: "
                "source membership cannot be verified."
            )
    result = report.get("result")
    if isinstance(result, dict) and "literature_evaluation" in result:
        from labcat.science.reporting import validated_literature_evaluation

        try:
            validated_literature_evaluation(result, report.get("sources", []))
        except (KeyError, TypeError, ValueError, AttributeError, OverflowError):
            raise ExportError(
                "The saved literature evaluation cannot be verified."
            ) from None
    identifier = report.get("id")
    if not isinstance(identifier, str) or not re.fullmatch(
        r"[A-Za-z0-9_-]{1,64}", identifier
    ):
        raise ExportError("The report has no supported export identity.")
    selected = {}
    for view in VIEWS[views]:
        if view == "sources":
            selected[view] = _source_records(report)
            continue
        text = report.get(VIEW_FIELDS[view])
        if not isinstance(text, str) or not text.strip():
            raise ExportError("The selected report section is unavailable.")
        selected[view] = _parsed(text)
        if isinstance(selected[view], str):
            lines = selected[view].splitlines()
            if lines and lines[0] in {"PI Summary:", "Technical Overview:"}:
                lines[0] = VIEW_LABELS[view] + ":"
                selected[view] = "\n".join(lines)
    if sum(len(_text(value)) for value in selected.values()) > MAX_REPORT_CHARACTERS:
        raise ExportError("The report exceeds the export size limit.")
    metadata = {"id": identifier}
    for key in ("title", "stage", "created_at"):
        value = report.get(key, "")
        if not isinstance(value, str) or len(value) > 300:
            raise ExportError("The report metadata cannot be exported.")
        metadata[key] = value
    return metadata, selected


def _text(value) -> str:
    return (
        value
        if isinstance(value, str)
        else json.dumps(value, indent=2, allow_nan=False)
    )


def _safe_source_url(url) -> bool:
    return approved_citation_url(url)


def _source_records(report: dict) -> list[dict]:
    """Export readable, report-owned source fields, never arbitrary
    annotations.

    Presentation may join several material records to one workspace
    source row. The Sources section retains that workspace row once,
    matching its saved ID.
    """
    sources = report.get("sources", [])
    if (
        not isinstance(sources, list)
        or len(sources) > 1000
        or any(not isinstance(source, dict) for source in sources)
    ):
        raise ExportError("The report source list is invalid.")
    records, seen = [], set()
    limits = {
        "id": 128,
        "title": 4096,
        "source_name": 300,
        "access_scope": 64,
        "provenance_status": 64,
        "kind": 64,
        "created_at": 100,
    }
    for source in sources:
        record = {}
        for field, limit in limits.items():
            value = source.get(field, "")
            if not isinstance(value, str) or len(value) > limit:
                raise ExportError("The report source details cannot be exported.")
            record[field] = " ".join(value.split())
        if record["id"]:
            if record["id"] in seen:
                continue
            seen.add(record["id"])
        record["title"] = _citation_title(record["title"]) or "Untitled source"
        record["source_name"] = record["source_name"] or "Source not recorded"
        record["url"] = (
            source["url"]
            if source.get("access_scope") == "public"
            and source.get("provenance_status") == "verified"
            and _safe_source_url(source.get("url"))
            else None
        )
        records.append(record)
    return records


def _source_details(source: dict) -> list[str]:
    lines = [
        source["source_name"],
        "Access: "
        + (source["access_scope"] or "not recorded")
        + " · Provenance: "
        + (source["provenance_status"] or "not recorded"),
    ]
    if source["kind"] == "discovery_reference":
        lines.append(
            "Public reference · review required. "
            "A listed reference alone does not establish material properties."
        )
    if source["created_at"]:
        lines.append("Saved: " + source["created_at"])
    if not source["url"]:
        lines.append("Approved public link unavailable.")
    return lines


def _source_text(records: list[dict]) -> str:
    if not records:
        return "No sources are saved with this report."
    lines = []
    for index, source in enumerate(records, 1):
        lines.extend([f"{index}. {source['title']}", *_source_details(source)])
        if source["url"]:
            lines.append(source["url"])
        lines.append("")
    return "\n".join(lines).rstrip()


def _links(report: dict, selected: dict) -> set[str]:
    """Only stored adapter citations become active links, never user-
    label URLs."""
    sources = report.get("sources", [])
    result = set()
    if isinstance(sources, list):
        for source in sources:
            if (
                isinstance(source, dict)
                and source.get("access_scope") == "public"
                and source.get("provenance_status") == "verified"
                and _safe_source_url(source.get("url"))
            ):
                result.add(source["url"])
    for value in selected.values():
        if isinstance(value, dict):
            result.update(_links({"sources": value.get("sources", [])}, {}))
    # CLI reports may carry only result provenance plus their saved text views.
    result_data = report.get("result")
    if isinstance(result_data, dict):
        candidates = (
            [
                *result_data.get("candidates", []),
                *result_data.get("comparison_records", []),
                *result_data.get("review_records", []),
            ]
            if all(
                isinstance(result_data.get(key, []), list)
                for key in ("candidates", "comparison_records", "review_records")
            )
            else []
        )
        if isinstance(candidates, list):
            for candidate in candidates:
                provenance = (
                    candidate.get("provenance", {})
                    if isinstance(candidate, dict)
                    else {}
                )
                if isinstance(provenance, dict):
                    for field in ("source_url", "method_reference"):
                        if _safe_source_url(provenance.get(field)):
                            result.add(provenance[field])
    return result


def _citation_links(report: dict) -> dict[str, str]:
    return {
        f"[{item['id']}]": item["url"]
        for item in report.get("references", [])
        if isinstance(item, dict)
        and re.fullmatch(r"[RS][1-9][0-9]{0,3}", str(item.get("id", "")))
        and _safe_source_url(item.get("url"))
    }


def _runs(text: str, links: set[str], citations: dict | None = None):
    """Return literal text and approved URL spans without interpreting
    markup."""
    offset = 0
    pattern = re.compile(URL_PATTERN.pattern + r"|\[[RS][1-9][0-9]{0,3}\]")
    for match in pattern.finditer(text):
        raw = match.group(0)
        if citations and raw in citations:
            yield text[offset : match.start()], None
            yield raw, citations[raw]
            offset = match.end()
            continue
        url = raw.rstrip(".,;:)]}")
        if url not in links:
            continue
        yield text[offset : match.start()], None
        yield url, url
        offset = match.start() + len(url)
    yield text[offset:], None


def _formula_pattern(report: dict):
    """Recognize only bounded, source-candidate formula labels for
    typography."""
    result = report.get("result")
    candidates = result.get("candidates", []) if isinstance(result, dict) else []
    comparisons = (
        result.get("comparison_records", []) if isinstance(result, dict) else []
    )
    reviews = result.get("review_records", []) if isinstance(result, dict) else []
    evaluation = (
        result.get("literature_evaluation", {}) if isinstance(result, dict) else {}
    )
    preliminary = (
        evaluation.get("ranked_candidates", [])
        if isinstance(evaluation, dict)
        and evaluation.get("version") == "literature-fit-v2"
        else []
    )
    if (
        not isinstance(candidates, list)
        or len(candidates) > 100
        or not isinstance(comparisons, list)
        or len(comparisons) > 36
        or not isinstance(reviews, list)
        or len(reviews) > 12
        or not isinstance(preliminary, list)
        or len(preliminary) > 12
    ):
        return None
    candidates = [
        *candidates,
        *comparisons,
        *reviews,
        *({"formula": row.get("name")} for row in preliminary if isinstance(row, dict)),
    ]
    formulas = set()
    for candidate in candidates:
        raw = candidate.get("formula") if isinstance(candidate, dict) else None
        if not isinstance(raw, str) or len(raw) > 160:
            continue
        formula = display_formula(raw)
        tokens = flat_formula_tokens(formula)
        if tokens and any(count for _, count in tokens):
            formulas.add(formula)
    if not formulas:
        return None
    # URL spans are always consumed literally, including unapproved URLs that
    # remain inert text. Never subscript a URL, DOI, record ID or longer token.
    names = "|".join(
        re.escape(name) for name in sorted(formulas, key=len, reverse=True)
    )
    return re.compile(
        URL_PATTERN.pattern
        + r"|(?<![\w.:/@+·\-])(?P<formula>"
        + names
        + r")(?![\w:/@+·\-]|\.[A-Za-z0-9_])"
    )


def _formula_runs(text: str, pattern):
    """Yield literal text and count spans; never parse general report
    markup."""
    if pattern is None:
        yield text, False
        return
    offset = 0
    for match in pattern.finditer(text):
        if match.group("formula") is None:
            continue
        yield text[offset : match.start()], False
        for element, count in flat_formula_tokens(match.group("formula")):
            yield element, False
            if count:
                yield count, True
        offset = match.end()
    yield text[offset:], False


def _formula_context(text: str, pattern):
    # S references are publication/background titles, not material labels.
    return None if re.match(r"^- \[S[1-9][0-9]{0,3}\]", text) else pattern


def _table_row(line: str):
    """Recognize literal, bounded pipe cells, without interpreting any
    markup."""
    stripped = line.strip()
    if not stripped.startswith("|") or not stripped.endswith("|"):
        return None
    cells = [cell.strip() for cell in stripped[1:-1].split("|")]
    if not 2 <= len(cells) <= MAX_TABLE_COLUMNS or any(
        len(cell) > MAX_TABLE_CELL_CHARACTERS for cell in cells
    ):
        return None
    return cells


def _blocks(value):
    """Only the application's small table/heading dialect is formatted.

    A whole malformed or oversized table remains literal text. No cells
    are discarded, HTML is never interpreted, and legacy JSON stays
    readable.
    """
    lines = _text(value).splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        header = _table_row(line)
        separator = _table_row(lines[index + 1]) if index + 1 < len(lines) else None
        if (
            isinstance(value, str)
            and header
            and separator
            and len(header) == len(separator)
            and all(len(cell) <= 120 for cell in header)
            and all(re.fullmatch(r":?-{3,}:?", cell) for cell in separator)
        ):
            end = index + 2
            rows = []
            valid = True
            while end < len(lines) and lines[end].strip().startswith("|"):
                row = _table_row(lines[end])
                valid = valid and row is not None and len(row) == len(header)
                rows.append(row)
                end += 1
            if valid and 0 < len(rows) <= MAX_TABLE_ROWS:
                yield "table", [header, *rows], 0
                index = end
                continue
        index += 1
        stripped = line.strip()
        if not stripped:
            yield "space", "", 0
        elif re.match(r"^\d+\.\s+", stripped):
            yield "candidate", stripped, 0
        elif (
            stripped.endswith(":")
            and len(stripped) < 100
            and not stripped.startswith(('"', "{", "["))
        ):
            yield "heading", stripped, 0
        elif stripped.startswith(('"', "{", "}", "[", "]")):
            yield "code", stripped, min(40, len(line) - len(line.lstrip()))
        else:
            yield "body", stripped, 0


def _column_weights(header: list[str]) -> list[float]:
    """Reserve space for prose while keeping rank and score columns
    compact."""
    if (
        len(header) == 6
        and header[:4] == ["Rank", "Material", "Screening priority", "Why considered"]
        and header[4] in {"Attribute evidence", "Relevant properties"}
        and header[5] == "Stability / key caveat"
    ):
        return [0.06, 0.13, 0.14, 0.27, 0.17, 0.23]
    if header in (
        [
            "Provisional rank",
            "Material",
            "Literature fit",
            "Selected criterion assessments",
            "Assessment coverage",
            "Stability & uncertainty",
        ],
        [
            "Fit rank",
            "Material",
            "Literature fit",
            "Criterion assessments",
            "Coverage",
            "Stability & uncertainty",
        ],
    ):
        return [0.13, 0.14, 0.15, 0.25, 0.14, 0.19]
    weights = []
    for cell in header:
        name = cell.casefold()
        if name == "supported score":
            weights.append(1.8)
        elif name in {"rank", "provisional rank", "score", "overall utility"}:
            weights.append(1.0 if name in {"rank", "provisional rank"} else 1.2)
        elif name in {"coverage", "evidence coverage", "assessment coverage"}:
            weights.append(1.5)
        elif name.startswith("material"):
            weights.append(1.8)
        elif name in {
            "selected evidence",
            "leading criteria",
            "selected criterion assessments",
        }:
            weights.append(2.8)
        elif name == "supported-field completeness":
            weights.append(2.8)
        else:
            weights.append(2.0)
    return [weight / sum(weights) for weight in weights]


def _ranking_header(header):
    return header[:3] in (
        ["Rank", "Material", "Overall utility"],
        ["Rank", "Material", "Supported score"],
        ["Rank", "Material (record)", "Overall utility"],
    )


def _export_blocks(value):
    """Repeat identity/utility across readable property groups in paged
    exports."""
    for kind, text, indent in _blocks(value):
        if kind == "table" and text[0] == [
            "Provisional rank",
            "Material",
            "Literature fit",
            "Selected criterion assessments",
            "Assessment coverage",
            "Stability & uncertainty",
        ]:
            # Short print labels preserve the same six fields without splitting
            # long header words in a portrait document.
            yield (
                "table",
                [
                    [
                        "Fit rank",
                        "Material",
                        "Literature fit",
                        "Criterion assessments",
                        "Coverage",
                        "Stability & uncertainty",
                    ],
                    *text[1:],
                ],
                indent,
            )
            continue
        if kind == "table" and _ranking_header(text[0]) and len(text[0]) > 7:
            for offset in range(3, len(text[0]), 4):
                last = min(offset + 1, len(text[0]) - 3)
                yield (
                    "heading",
                    f"Selected properties {offset - 2}–{last}:",
                    0,
                )
                yield "table", [row[:3] + row[offset : offset + 4] for row in text], 0
        else:
            yield kind, text, indent


def _score_fills(table: list, report: dict, view: str) -> dict[int, str]:
    """Color only a table row that matches the saved structured utility
    record."""
    result = report.get("result", {})
    evaluation = (
        result.get("literature_evaluation") if isinstance(result, dict) else None
    )
    if (
        len(table[0]) == 6
        and table[0][:4] == ["Rank", "Material", "Screening priority", "Why considered"]
        and table[0][4] in {"Attribute evidence", "Relevant properties"}
        and table[0][5] == "Stability / key caveat"
        and isinstance(evaluation, dict)
        and evaluation.get("version") == "literature-fit-v2"
    ):
        # Recreate the table from already revalidated structured evidence. A
        # merely similar historical/free-text table never receives score colors.
        lines = _preliminary_literature_lines(
            evaluation,
            report.get("sources", []),
            audit=table[0][4] == "Relevant properties",
            leads=result.get("candidate_leads", []),
            result=result,
        )
        expected = next(
            (text for kind, text, _ in _blocks("\n".join(lines)) if kind == "table"),
            None,
        )
        if table != expected:
            return {}
        return {
            index: screening_priority_presentation(row)["fill"]
            for index, row in enumerate(evaluation["ranked_candidates"], 1)
        }
    if not _ranking_header(table[0]):
        return {}
    contract = result.get("report_tables", {}) if isinstance(result, dict) else {}
    if not isinstance(contract, dict) or contract.get("schema") != "ranking-tables-v1":
        return {}
    table_record = contract.get("summary" if view == "pi" else "technical", {})
    rows = table_record.get("rows", []) if isinstance(table_record, dict) else []
    if not isinstance(rows, list) or len(rows) > MAX_TABLE_ROWS:
        return {}
    fills = {}
    for index, cells in enumerate(table[1:], 1):
        for row in rows:
            if not isinstance(row, dict):
                continue
            fill = score_color(row.get("score"))
            if fill and cells[:3] == [
                str(row.get("rank")),
                row.get("material"),
                f"{row['score']:.4f}",
            ]:
                analysis = row.get("score_analysis", {})
                if isinstance(analysis, dict) and (
                    analysis.get("status") == "needs_evidence"
                    or analysis.get("coverage", 1) < 1 - 1e-7
                ):
                    fill = "#E8ECEE"
                fills[index] = fill
                break
    return fills


def _passage_keep_next(blocks: list) -> set[int]:
    """Keep a short, complete passage provenance block on one Word page.

    Long excerpts and entire literature sections remain free to
    paginate. Only the exact five-line saved-report dialect forms this
    bounded layout group.
    """
    prefixes = (
        "- Source: ",
        "- Attributed excerpt (",
        "- Method context: ",
        "- Caveat: ",
        "- Retrieved-content SHA-256: ",
    )
    kept = set()
    for index in range(len(blocks) - len(prefixes) + 1):
        group = blocks[index : index + len(prefixes)]
        if (
            all(
                kind == "body" and text.startswith(prefix)
                for (kind, text, _), prefix in zip(group, prefixes, strict=True)
            )
            and sum(len(text) for _, text, _ in group) <= 2200
        ):
            kept.update(range(index, index + len(prefixes) - 1))
    return kept


def _pdf_fonts():
    """Embedded portable fonts; no dependence on a host desktop font
    install."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    with _FONT_LOCK:
        for name, filename in (
            ("LabcatSans", "DejaVuSans.ttf"),
            ("LabcatBold", "DejaVuSans-Bold.ttf"),
            ("LabcatMono", "DejaVuSansMono.ttf"),
            ("LabcatSerif", "DejaVuSerif.ttf"),
            ("LabcatSerifBold", "DejaVuSerif-Bold.ttf"),
        ):
            if name not in pdfmetrics.getRegisteredFontNames():
                path = files("labcat").joinpath("fonts", filename)
                pdfmetrics.registerFont(TTFont(name, str(path)))


def _pdf(
    metadata: dict,
    selected: dict,
    links: set[str],
    layout: ReportLayout,
    report: dict | None = None,
) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, letter
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import (
            CondPageBreak,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError:
        raise ExportUnavailable(
            "PDF support is not installed in this deployment."
        ) from None

    _pdf_fonts()
    stream = io.BytesIO()
    accent, tint, stripe = ACCENTS[layout.accent]
    teal = colors.HexColor(accent)
    paper = letter if layout.page_size == "letter" else A4
    width = paper[0] - 1.5 * inch
    regular = "LabcatSerif" if layout.font_family == "serif" else "LabcatSans"
    bold = "LabcatSerifBold" if layout.font_family == "serif" else "LabcatBold"
    leading = layout.font_size * (1.2 if layout.line_spacing == "compact" else 1.4)
    styles = {
        "body": ParagraphStyle(
            "Body",
            fontName="LabcatSans",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#223637"),
            spaceAfter=5,
        ),
        "heading": ParagraphStyle(
            "Heading",
            fontName="LabcatBold",
            fontSize=11,
            leading=15,
            textColor=teal,
            spaceBefore=8,
            spaceAfter=5,
            keepWithNext=True,
        ),
        "candidate": ParagraphStyle(
            "Candidate",
            fontName="LabcatBold",
            fontSize=11,
            leading=15,
            textColor=teal,
            spaceBefore=9,
            spaceAfter=5,
            keepWithNext=True,
        ),
        "code": ParagraphStyle(
            "AuditData",
            fontName="LabcatMono",
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#334747"),
            spaceAfter=2,
        ),
        "title": ParagraphStyle(
            "Title",
            fontName="LabcatBold",
            fontSize=23,
            leading=28,
            textColor=teal,
            spaceAfter=10,
            keepWithNext=True,
        ),
        "cell": ParagraphStyle(
            "TableCell",
            fontName="LabcatSans",
            fontSize=8.5,
            leading=11.5,
            textColor=colors.HexColor("#223637"),
            splitLongWords=True,
        ),
        "cell_heading": ParagraphStyle(
            "TableHeader",
            fontName="LabcatBold",
            fontSize=8.5,
            leading=11.5,
            textColor=teal,
            splitLongWords=True,
        ),
    }
    for name in ("body", "cell"):
        styles[name].fontName = regular
    for name in ("heading", "candidate", "title", "cell_heading"):
        styles[name].fontName = bold
    styles["body"].fontSize = layout.font_size
    styles["body"].leading = leading
    for name in ("heading", "candidate"):
        styles[name].fontSize = layout.font_size + 1
        styles[name].leading = leading + 1
    for name in ("cell", "cell_heading"):
        styles[name].fontSize = layout.font_size - 1.5
        styles[name].leading = leading - 1

    formula_pattern = _formula_pattern(report or {})

    def paragraph(text, style, *, formulas=True, literal_only=False, link_url=None):
        def literal(part):
            return "".join(
                f"<sub>{escape(value)}</sub>" if subscript else escape(value)
                for value, subscript in _formula_runs(
                    part, _formula_context(text, formula_pattern) if formulas else None
                )
            )

        spans = (
            [(text, link_url if link_url in links else None)]
            if literal_only
            else _runs(text, links, _citation_links(report or {}))
        )
        markup = "".join(
            (
                f'<link href="{escape(url, quote=True)}" color="{accent}">'
                f"{escape(part)}</link>"
                if url
                else literal(part)
            )
            for part, url in spans
        )
        return Paragraph(markup, style)

    if (report or {}).get("presentation_version") == PRESENTATION_VERSION:
        story = [
            paragraph("Labcat · Research report", styles["heading"]),
            paragraph(
                metadata["title"] or "Materials research",
                styles["title"],
                formulas=False,
            ),
        ]
        stamp = metadata["created_at"][:10] or "Date not recorded"
        mode = (
            "Saved report format"
            if report["format_source"] == "saved"
            else "Current Report Format"
        )
        story.extend([paragraph(stamp + " · " + mode, styles["body"]), Spacer(1, 8)])
    else:
        story = [paragraph("Labcat", styles["title"])]
        story.append(
            paragraph(metadata["title"] or "Saved research report", styles["heading"])
        )
        table = Table(
            [
                [
                    paragraph("SAVED REPORT", styles["heading"]),
                    paragraph("EXPORT", styles["heading"]),
                ],
                [
                    paragraph(metadata["id"], styles["body"]),
                    paragraph(
                        " + ".join(VIEW_LABELS[view] for view in selected),
                        styles["body"],
                    ),
                ],
                [
                    paragraph(
                        "Stage: " + (metadata["stage"] or "recorded"), styles["body"]
                    ),
                    paragraph(
                        metadata["created_at"] or "Date not recorded", styles["body"]
                    ),
                ],
            ],
            colWidths=[width / 2, width / 2],
            hAlign="LEFT",
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(stripe)),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 10),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]
            )
        )
        story.extend([table, Spacer(1, 14)])
    for index, (view, value) in enumerate(selected.items()):
        if index:
            story.append(PageBreak())
        story.append(paragraph(VIEW_LABELS[view], styles["title"]))
        if view == "sources":
            if not value:
                story.append(paragraph(_source_text(value), styles["body"]))
            for source_index, source in enumerate(value, 1):
                story.append(
                    paragraph(
                        f"{source_index}. {source['title']}",
                        styles["heading"],
                        formulas=False,
                        literal_only=True,
                        link_url=source["url"],
                    )
                )
                for line in _source_details(source):
                    story.append(
                        paragraph(
                            line, styles["body"], formulas=False, literal_only=True
                        )
                    )
            continue
        blocks = [block for block in _export_blocks(value) if block[0] != "space"]
        table_heading = None
        for block_index, (kind, text, indent) in enumerate(blocks):
            if block_index == 0 and text == VIEW_LABELS[view] + ":":
                continue
            if (
                kind == "heading"
                and block_index + 1 < len(blocks)
                and blocks[block_index + 1][0] == "table"
            ):
                # A normal keep-with-next would attempt to retain the complete
                # table. Keep just its heading, header and first row together.
                table_heading = paragraph(
                    text,
                    ParagraphStyle(
                        "TableHeading", parent=styles["heading"], keepWithNext=False
                    ),
                )
                continue
            if kind == "table":
                rows = [
                    [
                        paragraph(
                            cell, styles["cell_heading" if row_index == 0 else "cell"]
                        )
                        for cell in row
                    ]
                    for row_index, row in enumerate(text)
                ]
                grid = Table(
                    rows,
                    colWidths=[width * weight for weight in _column_weights(text[0])],
                    repeatRows=1,
                    splitByRow=1,
                    splitInRow=1,
                    hAlign="LEFT",
                )
                grid.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(tint)),
                            (
                                "ROWBACKGROUNDS",
                                (0, 1),
                                (-1, -1),
                                (
                                    [colors.white, colors.HexColor(stripe)]
                                    if layout.table_style == "striped"
                                    else [colors.white]
                                ),
                            ),
                            (
                                "LINEBELOW",
                                (0, 0),
                                (-1, 0),
                                0.6,
                                colors.HexColor("#BCCDC3"),
                            ),
                            (
                                "LINEBELOW",
                                (0, 1),
                                (-1, -1),
                                0.3,
                                colors.HexColor("#DCE5DF"),
                            ),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 6),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                            ("TOPPADDING", (0, 0), (-1, -1), 7),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                        ]
                    )
                )
                if layout.table_style == "grid":
                    grid.setStyle(
                        TableStyle(
                            [
                                (
                                    "GRID",
                                    (0, 0),
                                    (-1, -1),
                                    0.4,
                                    colors.HexColor("#CAD2D5"),
                                )
                            ]
                        )
                    )
                if layout.table_style == "minimal":
                    grid.setStyle(
                        TableStyle(
                            [
                                ("BACKGROUND", (0, 0), (-1, 0), colors.white),
                                ("LINEBELOW", (0, 1), (-1, -1), 0, colors.white),
                            ]
                        )
                    )
                for row_index, fill in _score_fills(text, report or {}, view).items():
                    grid.setStyle(
                        TableStyle(
                            [
                                (
                                    "BACKGROUND",
                                    (0, row_index),
                                    (0, row_index),
                                    colors.HexColor(fill),
                                ),
                                (
                                    "BACKGROUND",
                                    (2, row_index),
                                    (2, row_index),
                                    colors.HexColor(fill),
                                ),
                            ]
                        )
                    )
                if table_heading is not None:
                    grid.wrap(width, paper[1] - inch)
                    height = sum(grid._rowHeights[:2])
                    heading_height = table_heading.wrap(width, paper[1] - inch)[1]
                    story.extend(
                        [CondPageBreak(height + heading_height + 16), table_heading]
                    )
                    table_heading = None
                story.extend([grid, Spacer(1, 9)])
                continue
            style = styles[kind]
            if indent or (kind == "code" and text.endswith(("{", "["))):
                style = ParagraphStyle(
                    "AuditLine",
                    parent=style,
                    leftIndent=indent * 2,
                    keepWithNext=kind == "code" and text.endswith(("{", "[")),
                )
            story.append(paragraph(text, style))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#CAD9D2"))
        canvas.line(0.75 * inch, 0.58 * inch, paper[0] - 0.75 * inch, 0.58 * inch)
        canvas.setFont(regular, 8)
        canvas.setFillColor(teal)
        canvas.drawString(
            0.75 * inch, 0.4 * inch, "Saved report | citations and caveats retained"
        )
        if layout.page_numbers:
            canvas.drawRightString(
                paper[0] - 0.75 * inch, 0.4 * inch, str(document.page)
            )
        canvas.restoreState()

    document = SimpleDocTemplate(
        stream,
        pagesize=paper,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.8 * inch,
        title="Labcat saved research report",
        author="Labcat",
    )
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()


def _docx(
    metadata: dict,
    selected: dict,
    links: set[str],
    layout: ReportLayout,
    report: dict | None = None,
) -> bytes:
    try:
        from docx import Document
        from docx.opc.constants import RELATIONSHIP_TYPE
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        raise ExportUnavailable(
            "Word support is not installed in this deployment."
        ) from None

    document = Document()
    section = document.sections[0]
    page_width, page_height = (
        (8.5, 11) if layout.page_size == "letter" else (8.2677, 11.6929)
    )
    section.page_width, section.page_height = Inches(page_width), Inches(page_height)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    font_name = "Georgia" if layout.font_family == "serif" else "Calibri"
    accent, tint, stripe = ACCENTS[layout.accent]
    normal.font.name = font_name
    normal.font.size = Pt(layout.font_size)
    normal.paragraph_format.line_spacing = (
        1.2 if layout.line_spacing == "compact" else 1.4
    )
    normal.paragraph_format.space_after = Pt(5)
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        document.styles[name].font.color.rgb = RGBColor.from_string("000000")
        document.styles[name].font.name = font_name
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        fonts = document.styles[name].element.get_or_add_rPr().find(qn("w:rFonts"))
        if fonts is not None:
            # A template theme can otherwise override the selected serif face.
            for attribute in list(fonts.attrib):
                if attribute.rsplit("}", 1)[-1].lower().endswith("theme"):
                    del fonts.attrib[attribute]
    for name in ("Title", "Subtitle"):
        properties = document.styles[name].element.get_or_add_pPr()
        for border in properties.findall(qn("w:pBdr")):
            properties.remove(border)
    document.core_properties.title = "Labcat saved research report"
    document.core_properties.author = "Labcat"
    document.core_properties.comments = (
        "Exported from a persisted report without re-running research."
    )
    if (report or {}).get("presentation_version") == PRESENTATION_VERSION:
        document.add_paragraph("Labcat · Research report", "Subtitle")
        document.add_heading(metadata["title"] or "Materials research", 0)
        stamp = metadata["created_at"][:10] or "Date not recorded"
        mode = (
            "Saved report format"
            if report["format_source"] == "saved"
            else "Current Report Format"
        )
        document.add_paragraph(stamp + " · " + mode)
        for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
            document.styles[name].font.color.rgb = RGBColor.from_string(
                accent.lstrip("#")
            )
    else:
        document.add_heading("Materials research report", 0)
        document.add_paragraph("Labcat", "Subtitle")
        table = document.add_table(rows=3, cols=2)
        table.style = "Normal Table"
        for cells, values in zip(
            table.rows,
            [
                ("Saved report", metadata["id"]),
                ("Stage", metadata["stage"] or "recorded"),
                ("Created", metadata["created_at"] or "Date not recorded"),
            ],
            strict=True,
        ):
            for cell, value in zip(cells.cells, values, strict=True):
                cell.text = value
                properties = cell._tc.get_or_add_tcPr()
                shading = OxmlElement("w:shd")
                shading.set(
                    qn("w:fill"),
                    ("FFFFFF" if layout.table_style == "minimal" else tint.lstrip("#")),
                )
                properties.append(shading)
                if layout.table_style == "grid":
                    borders = OxmlElement("w:tcBorders")
                    for edge in ("top", "left", "bottom", "right"):
                        border = OxmlElement("w:" + edge)
                        border.set(qn("w:val"), "single")
                        border.set(qn("w:sz"), "4")
                        border.set(qn("w:color"), "CAD2D5")
                        borders.append(border)
                    properties.append(borders)

    formula_pattern = _formula_pattern(report or {})

    def append_text(paragraph, text, *, literal_only=False, link_url=None):
        spans = (
            [(text, link_url if link_url in links else None)]
            if literal_only
            else _runs(text, links, _citation_links(report or {}))
        )
        for part, url in spans:
            if not url:
                for value, subscript in _formula_runs(
                    part,
                    None if literal_only else _formula_context(text, formula_pattern),
                ):
                    run = paragraph.add_run(value)
                    if subscript:
                        run.font.subscript = True
                continue
            relation = paragraph.part.relate_to(
                url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
            )
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), relation)
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), accent.lstrip("#"))
            properties.append(color)
            run.append(properties)
            text_element = OxmlElement("w:t")
            text_element.text = part
            run.append(text_element)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)

    for index, (view, value) in enumerate(selected.items()):
        if index:
            document.add_page_break()
        document.add_heading(VIEW_LABELS[view], 1)
        if view == "sources":
            if not value:
                document.add_paragraph(_source_text(value))
            for source_index, source in enumerate(value, 1):
                append_text(
                    document.add_paragraph(style="Heading 2"),
                    f"{source_index}. {source['title']}",
                    literal_only=True,
                    link_url=source["url"],
                )
                for line in _source_details(source):
                    append_text(document.add_paragraph(), line, literal_only=True)
            continue
        blocks = list(_export_blocks(value))
        passage_keep_next = _passage_keep_next(blocks)
        for block_index, (kind, text, indent) in enumerate(blocks):
            if block_index == 0 and text == VIEW_LABELS[view] + ":":
                continue
            if kind == "space":
                continue
            if kind == "table":
                score_fills = _score_fills(text, report or {}, view)
                grid = document.add_table(rows=1, cols=len(text[0]))
                grid.style = "Normal Table"
                grid.autofit = False
                widths = [
                    Inches((page_width - 1.5) * weight)
                    for weight in _column_weights(text[0])
                ]
                for column, width in zip(grid.columns, widths, strict=True):
                    column.width = width
                header = OxmlElement("w:tblHeader")
                header.set(qn("w:val"), "true")
                grid.rows[0]._tr.get_or_add_trPr().append(header)
                for row_index, values in enumerate(text):
                    row = grid.rows[0] if row_index == 0 else grid.add_row()
                    if sum(map(len, values)) <= 1400 and max(map(len, values)) <= 400:
                        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                    for column_index, (cell, content, width) in enumerate(
                        zip(row.cells, values, widths, strict=True)
                    ):
                        cell.width = width
                        properties = cell._tc.get_or_add_tcPr()
                        shading = OxmlElement("w:shd")
                        fill = (
                            tint
                            if row_index == 0
                            else (
                                stripe
                                if layout.table_style == "striped"
                                and row_index % 2 == 0
                                else "#FFFFFF"
                            )
                        )
                        score_fill = (
                            score_fills.get(row_index)
                            if column_index in {0, 2}
                            else None
                        )
                        shading.set(
                            qn("w:fill"),
                            (
                                score_fill
                                or (
                                    "#FFFFFF"
                                    if layout.table_style == "minimal"
                                    else fill
                                )
                            ).lstrip("#"),
                        )
                        properties.append(shading)
                        borders = OxmlElement("w:tcBorders")
                        for edge in ("top", "left", "bottom", "right"):
                            border = OxmlElement("w:" + edge)
                            visible = (
                                layout.table_style == "grid"
                                or edge == "bottom"
                                and (row_index == 0 or layout.table_style == "striped")
                            )
                            border.set(qn("w:val"), "single" if visible else "nil")
                            border.set(qn("w:sz"), "4")
                            border.set(qn("w:color"), "CAD2D5")
                            borders.append(border)
                        properties.append(borders)
                        paragraph = cell.paragraphs[0]
                        paragraph.paragraph_format.space_after = Pt(4)
                        paragraph.paragraph_format.space_before = Pt(4)
                        paragraph.paragraph_format.keep_with_next = row_index == 0
                        append_text(paragraph, content)
                        for run in paragraph.runs:
                            run.font.size = Pt(layout.font_size - 1.5)
                            run.font.name = font_name
                            run.font.color.rgb = RGBColor.from_string(
                                accent.lstrip("#") if row_index == 0 else "223637"
                            )
                            run.bold = row_index == 0
                document.add_paragraph().paragraph_format.space_after = Pt(2)
                continue
            paragraph = document.add_paragraph(
                style="Heading 3" if kind in {"candidate", "heading"} else "Normal"
            )
            if block_index in passage_keep_next:
                paragraph.paragraph_format.keep_with_next = True
            if kind == "code":
                paragraph.paragraph_format.left_indent = Pt(indent * 2)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.keep_with_next = text.endswith(("{", "["))
            append_text(paragraph, text)
            if kind == "code":
                for run in paragraph.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(8)
    section.footer.paragraphs[0].text = (
        "Labcat | Saved report; citations and caveats retained"
    )
    if layout.page_numbers:
        footer = section.footer.paragraphs[0]
        footer.add_run(" | Page ")
        number = OxmlElement("w:fldSimple")
        number.set(qn("w:instr"), "PAGE")
        footer._p.append(number)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def render_download(
    report: dict,
    format: str,
    views: str = "both",
    *,
    fallback_layout: ReportLayout | None = None,
) -> tuple[bytes, str, str]:
    """Return bytes, content type and a title-independent safe
    attachment filename."""
    if format not in FORMATS:
        raise ExportError("Choose text, JSON, PDF or Word format.")
    if "references" not in report and isinstance(report.get("result"), dict):
        report = {
            **report,
            "references": citation_references(
                report["result"], report.get("sources", [])
            ),
        }
    metadata, selected = _content(report, views)
    layout, _ = report_layout(report, fallback_layout)
    extension = "txt" if format == "text" else format
    filename = f"labcat-{metadata['id']}-{views}.{extension}"
    if format == "json":
        document = {"report": metadata, "views": selected}
        # New saved views are prose for every presentation setting. Preserve the
        # machine-readable research in technical downloads, without widening PI
        # exports or copying arbitrary top-level workspace/connection metadata.
        if "audit" in selected and isinstance(report.get("result"), dict):
            document["result"] = report.get("original_result", report["result"])
        if report.get("presentation_version") == PRESENTATION_VERSION:
            document["presentation"] = {
                "version": PRESENTATION_VERSION,
                "format_source": report["format_source"],
                "settings_source": report["settings_source"],
                "settings": report["presentation"],
            }
            document["references"] = report["references"]
            document["presentation"]["material_names"] = report.get(
                "material_names", []
            )
            document["archive"] = {
                VIEW_FIELDS[view]: report["archive"][VIEW_FIELDS[view]]
                for view in selected
                if view in VIEW_FIELDS
            }
        try:
            payload = json.dumps(document, indent=layout.json_indent, allow_nan=False)
        except (TypeError, ValueError, RecursionError):
            raise ExportError("The saved report contains unsupported data.") from None
        if len(payload) > MAX_REPORT_CHARACTERS:
            raise ExportError("The report exceeds the export size limit.")
        return (payload + "\n").encode(), "application/json", filename
    if format == "text":
        parts = ["Labcat", metadata["title"]]
        if report.get("presentation_version") != PRESENTATION_VERSION:
            parts.append("Report: " + metadata["id"])
        for view, value in selected.items():
            lines = []
            content = _source_text(value) if view == "sources" else _text(value)
            for line in content.splitlines():
                lines.append(
                    line
                    if line.strip().startswith("|")
                    else textwrap.fill(
                        line,
                        width=layout.text_width,
                        break_long_words=False,
                        break_on_hyphens=False,
                        replace_whitespace=False,
                    )
                )
            parts.extend(["", VIEW_LABELS[view], "", "\n".join(lines)])
        if report.get("references") and any(view in selected for view in VIEW_FIELDS):
            parts.extend(["", "Source links", ""])
            for reference in report["references"]:
                parts.extend(
                    [f"[{reference['id']}] {reference['title']}", reference["url"], ""]
                )
        return (
            ("\n".join(parts).rstrip() + "\n").encode(),
            "text/plain; charset=utf-8",
            filename,
        )
    links = _links(report, selected)
    if format == "pdf":
        return (
            _pdf(metadata, selected, links, layout, report),
            "application/pdf",
            filename,
        )
    return (
        _docx(metadata, selected, links, layout, report),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename,
    )


def prepare_presentation(
    report: dict, format_source="saved", config=None, *, views="both"
) -> dict:
    """Derive a presentation without changing the saved record or
    scientific scores.

    Old workspace source rows retain public verification and URL scope
    while complete adapter provenance lives in result.candidates,
    comparison_records and review_records. Rejoin only those existing
    records, by exact source URL; no adapter, model or ranking call
    runs.
    """
    if format_source not in {"saved", "current"}:
        raise ExportError("Choose saved or current report formatting.")
    _content(report, views)
    original = report.get("result")
    result = deepcopy(original) if isinstance(original, dict) else {}
    execution = result.get("execution", {})
    saved = execution.get("presentation") if isinstance(execution, dict) else None
    try:
        if format_source == "saved" and isinstance(saved, dict):
            values = deepcopy(saved)
            settings_source = "saved-report"
            if "layout" not in values:
                values["layout"] = (config or load_config()).to_dict()["presentation"][
                    "layout"
                ]
                settings_source = (
                    "current-settings" if config is not None else "defaults"
                )
            appearance = config_from_mapping({"presentation": values})
        elif config is not None:
            appearance = config_from_mapping(
                {"presentation": config.to_dict()["presentation"]}
            )
            settings_source = "current-settings"
        else:
            appearance = load_config()
            settings_source = "defaults"
    except (ValueError, TypeError, RecursionError):
        raise ExportError("The report presentation preferences are invalid.") from None
    presentation = appearance.to_dict()["presentation"]
    sources = report.get("sources", [])
    if (
        not isinstance(sources, list)
        or len(sources) > 1000
        or any(not isinstance(source, dict) for source in sources)
    ):
        raise ExportError("The report source list is invalid.")
    from labcat.science.report_sources import report_scoped_sources

    try:
        sources = report_scoped_sources(result, sources)
    except (KeyError, TypeError, ValueError, AttributeError, OverflowError):
        raise ExportError("The report source snapshot cannot be verified.") from None
    candidates = result.get("candidates")
    legacy = not isinstance(candidates, list)
    if not legacy:
        if len(candidates) > 100:
            raise ExportError("The report shortlist exceeds its display limit.")
        comparison_records = result.get("comparison_records", [])
        if not isinstance(comparison_records, list) or len(comparison_records) > 36:
            raise ExportError(
                "The report comparison evidence exceeds its display limit."
            )
        review_records = result.get("review_records", [])
        if not isinstance(review_records, list) or len(review_records) > 12:
            raise ExportError("The report review evidence exceeds its display limit.")
        verified = {
            source.get("url"): source
            for source in sources
            if isinstance(source, dict)
            and source.get("access_scope") == "public"
            and source.get("provenance_status") == "verified"
            and source.get("kind") != "discovery_reference"
            and _safe_source_url(source.get("url"))
        }
        joined = []
        try:
            identities = {}
            for candidate in [*candidates, *comparison_records, *review_records]:
                provenance = candidate["provenance"]
                source = verified.get(provenance["source_url"])
                if source is None:
                    raise ValueError("Missing saved source")
                identifier = "material:" + candidate["material_id"]
                candidate["source_ids"] = [identifier]
                if identifier in identities:
                    if identities[identifier] != provenance:
                        raise ValueError("Conflicting saved source identity")
                    continue
                identities[identifier] = provenance
                joined.append(
                    {
                        **source,
                        "source_id": identifier,
                        "record_id": candidate["material_id"],
                        "provenance": deepcopy(provenance),
                    }
                )
            used = {source["url"] for source in joined}
            joined_ids = {source.get("id") for source in joined}
            sources = joined + [
                source
                for source in sources
                if source.get("kind") == "discovery_reference"
                or source.get("url") not in used
                # Distinct linked legacy rows can share a source URL. Keep their
                # membership without using them to invent another property record.
                or (source.get("id") is not None and source["id"] not in joined_ids)
            ]
            summary, technical = render_reports(result, appearance, sources)
            references = citation_references(result, sources)
        except (KeyError, TypeError, ValueError, AttributeError, OverflowError):
            raise ExportError(
                "The saved scientific report cannot be formatted safely."
            ) from None
    else:
        # Unknown older record schemas remain readable; never fabricate a modern
        # table from free-form prose or reinterpret an old model answer as facts.
        summary, technical = (
            report.get("pi_summary", ""),
            report.get("technical_audit", ""),
        )
        references = []
    result["execution"] = {
        **(execution if isinstance(execution, dict) else {}),
        "presentation": presentation,
    }
    prepared = {
        **report,
        "pi_summary": summary,
        "technical_audit": technical,
        "result": result,
        "sources": sources,
        "references": references,
        "presentation_version": PRESENTATION_VERSION,
        "format_source": format_source,
        "settings_source": settings_source,
        "presentation": presentation,
        "legacy": legacy,
        "archive": {
            "pi_summary": report.get("pi_summary", ""),
            "technical_audit": report.get("technical_audit", ""),
        },
        "original_result": deepcopy(original),
    }
    _content(prepared, views)
    prepared["material_names"] = offline_names(prepared)
    return prepared


def presentation_response(report: dict, chat_id: str) -> dict:
    """Only display fields and bounded preferences leave the
    representation API."""
    return {
        "version": PRESENTATION_VERSION,
        "chat_id": chat_id,
        "report_id": report["id"],
        "format_source": report["format_source"],
        "settings_source": report["settings_source"],
        "presentation": report["presentation"],
        "pi_summary": report["pi_summary"],
        "technical_audit": report["technical_audit"],
        "report_tables": (
            report["result"].get("report_tables") if not report["legacy"] else None
        ),
        "references": report["references"],
        "legacy": report["legacy"],
        "material_names": report.get("material_names", []),
    }


def create_exports_router(store, config_reader=None, name_source_reader=None):
    """Presentation and downloads share scoped, read-only saved-data
    rendering."""
    from fastapi import APIRouter, HTTPException, Response

    router = APIRouter(prefix="/api/chats")
    name_resolver = ChemicalNameResolver()

    def name_source_options():
        if name_source_reader is None:
            return {"allow_literature": True, "allow_openalex": True}
        selected = name_source_reader()
        enabled = (
            selected["enabled_sources"] if selected["search_public_references"] else []
        )
        return {
            "allow_literature": "europe_pmc" in enabled,
            "allow_openalex": "openalex" in enabled,
        }

    def read(chat_id, report_id, format_source, views="both"):
        try:
            detail = store.get_global_chat(chat_id)
            report = next(
                (item for item in detail["reports"] if item["id"] == report_id), None
            )
            if report is None:
                raise WorkspaceNotFound("Report not found in this chat.")
            report = {
                **report,
                "sources": [
                    source
                    for source in detail["sources"]
                    if source["id"] in report["source_ids"]
                ],
            }
            stored_result = report.get("result")
            execution = (
                stored_result.get("execution", {})
                if isinstance(stored_result, dict)
                else {}
            )
            if not isinstance(execution, dict):
                execution = {}
            has_saved = (
                isinstance(execution.get("presentation"), dict)
                and "layout" in execution["presentation"]
            )
            config = (
                config_reader()
                if config_reader and (format_source == "current" or not has_saved)
                else None
            )
            prepared = prepare_presentation(report, format_source, config, views=views)
            prepared["material_names"] = name_resolver.names(
                prepared, **name_source_options()
            )
            return prepared
        except WorkspaceNotFound:
            raise HTTPException(404, "Report not found in this chat.") from None
        except sqlite3.Error:
            raise HTTPException(503, "Saved reports are unavailable.") from None
        except ExportError as error:
            raise HTTPException(422, str(error)) from None

    @router.get("/{chat_id}/reports/{report_id}/presentation")
    def presentation(
        chat_id: str,
        report_id: str,
        format_source: Literal["saved", "current"] = "saved",
    ):
        report = read(chat_id, report_id, format_source)
        return presentation_response(report, chat_id)

    @router.post("/{chat_id}/reports/{report_id}/chemical-names")
    def chemical_names(
        chat_id: str, report_id: str, request: ChemicalNameRequest | None = None
    ):
        report = read(chat_id, report_id, "saved")
        return {
            "chat_id": chat_id,
            "report_id": report_id,
            "material_names": name_resolver.names(
                report, lookup=True, **name_source_options()
            ),
        }

    @router.get("/{chat_id}/reports/{report_id}/export")
    def export(
        chat_id: str,
        report_id: str,
        format: Literal["text", "json", "pdf", "docx"] = "text",
        views: Literal[
            "pi", "audit", "both", "sources", "pi,sources", "audit,sources", "all"
        ] = "both",
        format_source: Literal["saved", "current"] = "saved",
    ):
        report = read(chat_id, report_id, format_source, views)
        try:
            body, media_type, filename = render_download(report, format, views)
        except ExportUnavailable as error:
            raise HTTPException(503, str(error)) from None
        except ExportError as error:
            raise HTTPException(422, str(error)) from None
        return Response(
            body,
            media_type=media_type,
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "X-Labcat-Layout-Source": report["settings_source"],
                "Cache-Control": "no-store",
            },
        )

    return router


def _template_report(presentation: dict) -> dict:
    """Layout examples contain explicit placeholders, never
    demonstration results."""
    detail = presentation["verbosity"] != "concise"
    pi = (
        "Summary:\n\nTemplate preview — no research performed.\n\n"
        "Research objective:\n\nQuestion: [Your research question "
        "and selected priorities.]\n\n"
        "Written summary:\n\nReport: [The report explains "
        "the strongest supported candidates, "
        "the ranking rationale, and the most important uncertainties.]\n\n"
        "Shortlist:\n\n| Rank | Material | Supported score | "
        "Leading criteria | Key caveat |\n"
        "| --- | --- | --- | --- | --- |\n"
        "| [Rank] | [Supported material] | [Supported score] | "
        "[Source-backed rationale] | "
        "[Unresolved context] |\n\n"
        "Public evidence and caveats:\n\n"
        "Evidence: [Verified public citations accompany scientific claims. "
        "Missing values "
        "remain marked unknown; no scientific evidence is included in this preview.]"
    )
    audit = (
        "Technical View:\n\nTemplate preview — no research performed.\n\n"
        "Research context:\n\nScope: [Question, material class, application "
        "and ranking profile.]\n\n"
        "Expanded shortlist:\n\n"
        "| Rank | Material | Supported score | [Selected property] | "
        "[Another selected property] |\n"
        "| --- | --- | --- | --- | --- |\n"
        "| [Rank] | [Supported material] | [Supported score] | "
        "[Source value and unit, or Unknown] | [Source value or Unavailable] |\n\n"
        "Score meaning: Supported score sums known weighted contributions. "
        "Incomplete evidence is shown neutrally; missing properties do not "
        "demonstrate poor performance. Scores and coverage are not confidence "
        "or probability. "
        "All positively weighted properties remain visible. Wide PDF and Word "
        "tables repeat rank, identity, and supported score across property groups.\n\n"
        "Ranking rationale:\n\nExplanation: [Selected criteria, "
        "their relative importance, "
        "and the evidence contributing to each candidate's score.]\n\n"
        "Evidence and source traceability:\n\n"
        "Citations: [Short public-source references link to the supporting records. "
        "Raw identifiers, hashes and execution data remain in Audit details "
        "and the JSON archive.]\n\n"
        "Limitations:\n\nReview: [Missing data, phase or method differences "
        "and applicability "
        "caveats remain visible. This preview contains no research results.]"
    )
    if detail:
        audit += (
            "\n\nAttribute evidence follow-up:\n\n"
            "| Attribute | Search status | Evidence status |\n| --- | --- | --- |\n"
            "| [Missing selected attribute] | [Lookup outcome] | "
            "[Unscored review lead] |\n\n"
            "Review leads: [Attributable literature passages remain separate "
            "from verified "
            "candidate properties and never silently change the ranking.]"
        )
    if presentation["verbosity"] == "detailed":
        pi += (
            "\n\nNext review steps:\n\nFollow-up: [The report highlights the public "
            "evidence and unresolved context to inspect before "
            "interpreting the shortlist.]"
        )
    return {
        "id": "format-preview",
        "title": "Report format template",
        "stage": "template — no research performed",
        "created_at": "Preview only",
        "pi_summary": pi,
        "technical_audit": audit,
        "sources": [],
        "result": {
            "stage": "template",
            "candidates": [],
            "execution": {"presentation": presentation},
        },
    }


def _preview_css() -> str:
    """Finite classes only; the preview does not need inline CSS or
    remote assets."""
    css = (
        "*{box-sizing:border-box}body{margin:0;background:#edf1ef;color:#223637}"
        ".sheet{max-width:100%;margin:0 auto;padding:32px;background:white;"
        "overflow-wrap:anywhere}.paper-letter{width:816px}.paper-a4{width:794px}"
        ".font-sans{font-family:Arial,sans-serif}.font-serif{font-family:Georgia,serif}"
        ".size-10{font-size:13.333px}.size-11{font-size:14.667px}"
        ".size-12{font-size:16px}"
        ".spacing-compact{line-height:1.2}.spacing-comfortable{line-height:1.4}"
        "h1{font-size:1.8em;margin:0 0 16px}h2{font-size:1.5em;margin:24px 0 12px}"
        "h3{font-size:1.1em;margin:20px 0 8px}p{margin:0 0 10px}"
        ".notice{font-weight:bold}table{border-collapse:collapse;width:100%;"
        "table-layout:fixed;margin:12px 0 18px;font-size:.87em}"
        "td,th{text-align:left;vertical-align:top;padding:8px}"
        "th{border-bottom:1px solid #CAD2D5}"
        ".table-striped td{border-bottom:1px solid #CAD2D5}"
        ".table-grid td,.table-grid th{border:1px solid #CAD2D5}"
        ".table-minimal th{background:white!important}"
        "pre{white-space:pre-wrap;overflow-wrap:anywhere}"
        "footer{border-top:1px solid #CAD2D5;margin-top:24px;"
        "padding-top:12px;font-size:.8em}"
        "@media(max-width:500px){.sheet{padding:18px}td,th{padding:5px}}"
    )
    for name, (accent, tint, stripe) in ACCENTS.items():
        css += (
            f".accent-{name} h1,.accent-{name} h2,.accent-{name} h3"
            f"{{color:{accent}}}.accent-{name} th{{color:{accent};background:{tint}}}"
            f".accent-{name}.table-striped tbody tr:nth-child(even)"
            f"{{background:{stripe}}}"
        )
    return css


def _preview_html(report: dict, views: str) -> str:
    layout, _ = report_layout(report)
    _, selected = _content(report, views)
    classes = (
        f"sheet paper-{layout.page_size} font-{layout.font_family} "
        f"size-{layout.font_size} spacing-{layout.line_spacing} "
        f"accent-{layout.accent} table-{layout.table_style}"
    )
    html = [
        '<!doctype html><html lang="en"><head><meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width, initial-scale=1">',
        '<link rel="stylesheet" href="/api/report-preview/styles.css">',
        "<title>Report format template preview</title></head><body>",
        f'<main class="{classes}"><h1>Report format template</h1>',
        '<p class="notice">Template preview — no research performed.</p>',
    ]
    for view, value in selected.items():
        html.append(f"<section><h2>{VIEW_LABELS[view]}</h2>")
        for kind, text, _ in _blocks(value):
            if kind == "space" or text == VIEW_LABELS[view] + ":":
                continue
            if kind == "table":
                html.append("<table><thead><tr>")
                html.extend(f'<th scope="col">{escape(cell)}</th>' for cell in text[0])
                html.append("</tr></thead><tbody>")
                for row in text[1:]:
                    html.append(
                        "<tr>"
                        + "".join(f"<td>{escape(cell)}</td>" for cell in row)
                        + "</tr>"
                    )
                html.append("</tbody></table>")
            else:
                tag = (
                    "h3"
                    if kind in {"heading", "candidate"}
                    else "pre" if kind == "code" else "p"
                )
                html.append(f"<{tag}>{escape(text)}</{tag}>")
        html.append("</section>")
    html.append(
        "<footer>Template only; scientific evidence and caveats remain required."
    )
    if layout.page_numbers:
        html.append(" Page numbers are included in PDF and Word exports.")
    html.append("</footer></main></body></html>")
    return "".join(html)


def create_report_preview_router(config_reader=None):
    """Safe template previews; never research, write a workspace or
    accept content."""
    from fastapi import APIRouter, HTTPException, Response

    router = APIRouter(prefix="/api/report-preview")
    cache, cache_lock = OrderedDict(), threading.Lock()
    ttl, max_entries, max_bytes = 120, 8, 2_000_000

    def parsed(request, *, download=False):
        allowed = {"presentation", "views"} | ({"format"} if download else set())
        try:
            if not isinstance(request, dict) or request.keys() - allowed:
                raise ValueError(
                    "Provide only report presentation and preview options."
                )
            if len(json.dumps(request, allow_nan=False)) > 8000:
                raise ValueError("Preview preferences exceed the size limit.")
            base = config_reader() if config_reader else load_config()
            supplied = request.get("presentation", base.to_dict()["presentation"])
            if not isinstance(supplied, dict):
                raise ValueError("Invalid preview presentation.")
            presentation = {**base.to_dict()["presentation"], **supplied}
            if "layout" in supplied:
                presentation["layout"] = asdict(
                    validate_layout(supplied["layout"], base.presentation.layout)
                )
            validated = config_from_mapping({"presentation": presentation})
            normalized = validated.to_dict()["presentation"]
            views = request.get("views", "both")
            format = request.get("format", normalized["format"])
            if views not in VIEWS or format not in FORMATS:
                raise ValueError("Choose a supported preview format and report view.")
            return _template_report(normalized), normalized, views, format
        except sqlite3.Error:
            raise HTTPException(
                503, "Saved report preferences are unavailable."
            ) from None
        except (ValueError, TypeError, RecursionError):
            raise HTTPException(
                422, "Choose supported report presentation preferences only."
            ) from None

    def exported(request):
        report, _, views, format = parsed(request, download=True)
        try:
            body, media_type, filename = render_download(report, format, views)
            if len(body) > max_bytes:
                raise ExportError("The preview exceeds its download size limit.")
            return body, media_type, filename, format
        except ExportUnavailable as error:
            raise HTTPException(503, str(error)) from None
        except ExportError as error:
            raise HTTPException(422, str(error)) from None

    def attachment(body, media_type, filename):
        return Response(
            body,
            media_type=media_type,
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Cache-Control": "no-store",
                "X-Content-Type-Options": "nosniff",
            },
        )

    def remember(body, media_type, filename, format):
        if len(body) > max_bytes:
            raise HTTPException(422, "The preview exceeds its size limit.")
        token, now = secrets.token_hex(16), time.monotonic()
        with cache_lock:
            for key in list(cache):
                if cache[key][0] <= now:
                    del cache[key]
            while len(cache) >= max_entries:
                cache.popitem(last=False)
            cache[token] = (now + ttl, body, media_type, filename, format)
        return token

    def recalled(token, format):
        if not re.fullmatch(r"[0-9a-f]{32}", token):
            raise HTTPException(404, "Preview not found.")
        with cache_lock:
            saved = cache.get(token)
            if saved is None or saved[0] <= time.monotonic():
                cache.pop(token, None)
                raise HTTPException(410, "Preview expired; generate it again.")
            _, body, media_type, filename, selected_format = saved
        if selected_format != format:
            raise HTTPException(404, "Preview format does not match this request.")
        return body, media_type, filename

    @router.get("/styles.css")
    def styles():
        return Response(
            _preview_css(),
            media_type="text/css",
            headers={
                "X-Content-Type-Options": "nosniff",
                "Cache-Control": "no-store",
            },
        )

    @router.post("")
    def preview(request: dict):
        report, presentation, views, _ = parsed(request)
        html = _preview_html(report, views)
        token = remember(html.encode(), "text/html", "", "html")
        return {
            "label": "Template preview — no research performed",
            "html": html,
            "render_url": f"/api/report-preview/render/{token}",
            "documents": {
                "pi_summary": report["pi_summary"],
                "technical_audit": report["technical_audit"],
            },
            "text": render_download(report, "text", views)[0].decode(),
            "json": render_download(report, "json", views)[0].decode(),
            "presentation": presentation,
            "layout": presentation["layout"],
            "layout_note": "Representative Word layout; pagination may vary.",
            "download_endpoint": "/api/report-preview/download",
        }

    @router.post("/download")
    def download(request: dict):
        body, media_type, filename, _ = exported(request)
        return attachment(body, media_type, filename)

    @router.post("/download-link")
    def download_link(request: dict):
        body, media_type, filename, format = exported(request)
        token = remember(body, media_type, filename, format)
        return {
            "url": f"/api/report-preview/download/{token}?format={format}",
            "filename": filename,
        }

    @router.get("/download/{token}")
    def ticket(token: str, format: Literal["text", "json", "pdf", "docx"]):
        body, media_type, filename = recalled(token, format)
        return attachment(body, media_type, filename)

    @router.get("/render/{token}")
    def render_preview(token: str):
        body, media_type, _ = recalled(token, "html")
        return Response(
            body,
            media_type=media_type,
            headers={
                "Cache-Control": "no-store",
                "X-Content-Type-Options": "nosniff",
            },
        )

    return router
